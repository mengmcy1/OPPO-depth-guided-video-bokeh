"""Fill the existing next-report efficiency slide and narration from measured data."""
from copy import deepcopy
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree as ET

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "汇报/下次汇报草稿_基线更新与ATB分析"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
A = "{" + NS["a"] + "}"


def text(node):
    return "\n".join(t.text or "" for t in node.findall(".//a:t", NS))


def put(node, value, size=None):
    body = node.find("p:txBody", NS)
    if body is None:
        body = node.find("a:txBody", NS)
    assert body is not None
    template = body.find("a:p", NS)
    run = body.find(".//a:r", NS)
    for paragraph in list(body.findall("a:p", NS)):
        body.remove(paragraph)
    for line in str(value).split("\n"):
        paragraph = deepcopy(template)
        for child in list(paragraph):
            if child.tag != A + "pPr":
                paragraph.remove(child)
        replacement = deepcopy(run)
        for child in list(replacement):
            if child.tag != A + "rPr":
                replacement.remove(child)
        properties = replacement.find("a:rPr", NS)
        if properties is None:
            properties = ET.SubElement(replacement, A + "rPr")
        if size is not None:
            properties.set("sz", str(round(size * 100)))
        ET.SubElement(replacement, A + "t").text = line
        paragraph.append(replacement)
        body.append(paragraph)


def main():
    record = json.loads((ROOT / "03_实验记录/records/video-e2e-4090d-20261001.json").read_text())
    assert record["status"] == "SEVEN_METHOD_VIDEO_E2E_COMPLETE"
    rows = {row["method"]: row for row in record["table"]}
    path = REPORT / "基线更新与ATB效率分析_草稿.pptx"
    with ZipFile(path) as file:
        contents = {name: file.read(name) for name in file.namelist()}
    changed = {}
    found = False
    for name, data in contents.items():
        if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
            continue
        slide = ET.fromstring(data)
        shapes = slide.findall(".//p:sp", NS)
        titles = [text(shape) for shape in shapes]
        if "视频端到端效率" in titles:
            found = True
            table = slide.find(".//a:tbl", NS)
            assert table is not None
            grid = table.find("a:tblGrid", NS)
            table_rows = table.findall("a:tr", NS)
            assert len(table_rows) == 8
            if len(grid) == 5:
                grid.insert(2, deepcopy(grid[2]))
                for row in table_rows:
                    row.insert(2, deepcopy(row.findall("a:tc", NS)[2]))
            assert len(grid) == 6
            for column, width in zip(grid, [2.55, 0.9, 2.05, 2.05, 1.35, 2.3]):
                column.set("w", str(round(width * 914400)))
            values = [["方法", "设备", "Core (s/25帧) ↓", "E2E (s/25帧) ↓", "FPS ↑", "峰值内存 (GiB) ↓"]]
            names = {"deeplens": "DeepLens", "drbokeh": "Dr.Bokeh", "bokehme": "BokehMe",
                     "mpib": "MPIB", "bokehdiff": "BokehDiff", "magicbokeh": "MagicBokeh†", "any2bokeh": "Any-to-Bokeh"}
            for method in names:
                row = rows[method]
                memory = row["CPU_RSS_GiB_sampled"] if method == "deeplens" else row["Peak_device_memory_GiB_sampled"]
                values.append([names[method], "CPU" if method == "deeplens" else "GPU",
                    f'{row["Core_seconds_mean"]:.3f}', f'{row["E2E_given_inputs_seconds_mean"]:.3f}',
                    f'{row["FPS_E2E"]:.3f}', f'{memory:.2f}'])
            for row, fields in zip(table_rows, values):
                for cell, value in zip(row.findall("a:tc", NS), fields):
                    put(cell, value, 16 if row is table_rows[0] else 19)
            for shape in shapes:
                original = text(shape)
                if "当前服务器" in original:
                    put(shape, "RTX 4090 D · 1024×576、25帧 · 每输入预热1次、正式5次", 15)
                elif "从读取给定输入" in original:
                    put(shape, "从读取给定输入到25帧CPU uint8 RGB；不含模型加载、外部标注生成、写盘与编码。\n"
                        "†MagicBokeh包含自身估深、GT目标辅助选点与颜色校正；Core为估深网络及生成器GPU段。\n"
                        "GPU行报告设备采样显存；DeepLens为Xeon 8481C / TF1.15 / 8线程，内存列为CPU RSS。", 12.5)
            changed[name] = ET.tostring(slide, xml_declaration=True, encoding="UTF-8", standalone=True)
        elif "当前进度与下一步" in titles:
            for shape in shapes:
                original = text(shape)
                if "当前服务器的视频端到端效率尚待测量" in original:
                    put(shape, "六种GPU方法及DeepLens CPU完整25帧测量完成\n每方法20次正式重复；旧机记录分开保留")
                elif "先完成已确定的性能评价" in original:
                    put(shape, "编码分批、缓存优化和失败分析暂不开展\n本轮性能评价已完成，后续等待明确执行任务")
            changed[name] = ET.tostring(slide, xml_declaration=True, encoding="UTF-8", standalone=True)
    assert found
    with ZipFile(path, "w", ZIP_DEFLATED) as file:
        for name, data in contents.items():
            file.writestr(name, changed.get(name, data))
    document_path = REPORT / "基线更新与ATB效率分析_讲稿草稿.docx"
    document = Document(document_path)
    atb, magic, deep = [rows[method] for method in ["any2bokeh", "magicbokeh", "deeplens"]]
    replacements = {
        "这一页单独呈现当前服务器处理完整25帧视频":
            f'这一页呈现当前服务器的完整25帧实测。六种GPU方法使用RTX 4090 D，DeepLens使用Xeon Platinum 8481C的TF1.15 CPU路径。四个固定输入各预热一次、正式五次，七方法共140次正式测量。ATB的Core为{atb["Core_seconds_mean"]:.3f}秒，E2E为{atb["E2E_given_inputs_seconds_mean"]:.3f}秒，吞吐率为{atb["FPS_E2E"]:.3f}fps。视频时间来自实际完整25帧执行，不由模型单帧时间外推。',
        "DeepLens报告CPU进程内存，其余":
            f'DeepLens固定intra-op8线程、inter-op1线程；Core为{deep["Core_seconds_mean"]:.3f}秒，E2E为{deep["E2E_given_inputs_seconds_mean"]:.3f}秒，CPU RSS采样峰值为{deep["CPU_RSS_GiB_sampled"]:.3f}GiB。其余方法报告GPU设备采样显存，PyTorch allocated/reserved另存，CPU与GPU内存不能作同质排序。MagicBokeh本次包含每帧重新估深、GT目标辅助选点、条件构造/传输、六图块生成与拼接，以及PIL/AdaIN颜色处理；其Core为{magic["Core_seconds_mean"]:.3f}秒，E2E为{magic["E2E_given_inputs_seconds_mean"]:.3f}秒。Core包含估深网络forward和生成器GPU段，因此与上一页generator-only模型成本的边界不同。外部benchmark标注生成、模型加载、最终写盘、编码和指标计算不计。',
        "第一，七方法Synthetic、六方法DAVIS和新旧GPU的模型成本":
            '第一，七方法Synthetic、六方法DAVIS和新旧GPU模型成本结果继续保留。本轮只补当前服务器的七方法视频端到端效率，完成完整25帧、四输入、各预热一次和正式五次的测量。旧PRO 6000 Blackwell及Xeon 8470Q结果按原硬件标注保留；旧机MagicBokeh没有视频E2E记录，继续留空。',
        "第二，后续若补新卡端到端效率":
            '第二，当前服务器的视频端到端效率已完成；每次的Core、E2E、FPS、显存或RSS以及原始日志均已保存。均值、标准差和每K统计来自重复实测。新旧机器的GPU、CPU及环境存在差异，两张表服务于各自明确的成本边界，不将跨机差异解释为模型优化收益。'
    }
    for paragraph in document.paragraphs:
        for prefix, replacement in replacements.items():
            if paragraph.text.startswith(prefix):
                properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
                paragraph.clear()
                run = paragraph.add_run(replacement)
                if properties is not None:
                    run._r.insert(0, properties)
                break
    document.save(document_path)
    print("UPDATED_NEXT_REPORT_E2E_SLIDE_AND_NARRATION", list(changed))


if __name__ == "__main__":
    main()
